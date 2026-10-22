"""月考成绩册 Excel 导出。"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="EFF6FF")
HEADER_FONT = Font(bold=True, color="1E3A8A")
THIN = Side(style="thin", color="D8E0EC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _style_header(ws, row, columns):
    for col in range(1, columns + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def build_overview_workbook(data: dict) -> bytes:
    """把成绩总览结果转成可直接下发的 xlsx 字节。"""
    subjects = data["subjects"]
    has_totals = data["has_total_access"]
    headers = ["学生姓名", "班级", "年级"]
    if has_totals:
        headers += ["总分", "总分满分", "总分班内排名", "总分范围排名", "未录入科目"]
    for subject in subjects:
        headers += [f"{subject}得分", f"{subject}满分", f"{subject}班内排名", f"{subject}范围排名"]

    wb = Workbook()
    ws = wb.active
    ws.title = "月考成绩"
    last_col = len(headers)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    title = ws.cell(row=1, column=1, value=f"{data['exam_month'] or ''} 月考成绩册".strip())
    title.font = Font(bold=True, size=14)
    title.alignment = CENTER
    ws.row_dimensions[1].height = 26

    header_row = 2
    for index, name in enumerate(headers, 1):
        ws.cell(row=header_row, column=index, value=name)
    _style_header(ws, header_row, last_col)

    for row in data["items"]:
        if has_totals:
            can_view_total = row["can_view_total"]
            values = [
                row["student_name"], row["class_name"], row["grade"],
                row["total_score"] if can_view_total else "仅可见任教学科",
                row["total_max_score"] if can_view_total else None,
                row["class_rank"] if can_view_total else None,
                row["scope_rank"] if can_view_total else None,
                "、".join(row["missing_subjects"]) if can_view_total and row["missing_subjects"] else None,
            ]
        else:
            values = [row["student_name"], row["class_name"], row["grade"]]
        for subject in subjects:
            cell = row["scores"].get(subject)
            if cell:
                values += [cell["score"], cell["max_score"], cell["class_rank"], cell["scope_rank"]]
            else:
                values += [None, None, None, None]
        ws.append(values)

    last_row = max(ws.max_row, header_row)
    for row in ws.iter_rows(min_row=header_row, max_row=last_row, max_col=last_col):
        for cell in row:
            cell.border = BORDER
            cell.alignment = CENTER

    widths = {"A": 16, "B": 14, "C": 10}
    for index in range(1, last_col + 1):
        letter = get_column_letter(index)
        ws.column_dimensions[letter].width = widths.get(letter, 14)
    ws.freeze_panes = "D3"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(last_col)}{last_row}"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
