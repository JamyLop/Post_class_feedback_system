from datetime import date

import pytest

from app.core.security import hash_password
from app.models.class_ import Class, ClassStudent, ClassTeacher
from app.models.student_case import CaseCycle, StudentCase, SubjectPlan
from app.models.user import User
from app.models.weekly_score import WeeklyTestScore

API = "/api/monthly-exam-scores"


def initial_case(db, cls, student_id, subjects, cycle=None, status="draft"):
    if cycle is None:
        cycle = CaseCycle(name="成绩初评测试", grade=cls.grade, school_year="2026-2027",
                          starts_on=date(2026, 8, 1), ends_on=date(2027, 7, 31))
        db.add(cycle)
        db.flush()
    case = StudentCase(cycle_id=cycle.id, student_id=student_id, class_id=cls.id,
                       owner_teacher_id=cls.teacher_id, status=status)
    db.add(case)
    db.flush()
    for subject in subjects:
        db.add(SubjectPlan(student_case_id=case.id, subject=subject, teacher_id=cls.teacher_id))
    return case


@pytest.fixture()
def exam_roster(db, seed_users):
    password = hash_password("test123456")
    subject = User(username="exam_subject", name="任课老师", role="subject_teacher", password_hash=password)
    students = [User(username=f"exam_student_{i}", name=f"测试学生{i}", role="student", password_hash=password) for i in range(5)]
    db.add_all([subject, *students])
    db.flush()
    classes = [Class(name=name, grade=grade, teacher_id=seed_users[teacher]) for name, grade, teacher in (
        ("甲班", "高三", "teacher1"), ("乙班", "高三", "teacher2"),
        ("丙班", "高二", "teacher2"), ("丁班", "高三", "teacher2"))]
    db.add_all(classes)
    db.flush()
    a, b, c, d = classes
    db.add_all([
        ClassTeacher(class_id=a.id, teacher_id=subject.id, role="subject_teacher", subject="数学"),
        ClassTeacher(class_id=b.id, teacher_id=subject.id, role="subject_teacher", subject="英语"),
        ClassTeacher(class_id=a.id, teacher_id=seed_users["teacher2"], role="subject_teacher", subject="数学"),
    ])
    ids = [seed_users[f"student{i}"] for i in range(1, 4)] + [student.id for student in students]
    cycle = CaseCycle(name="月考初评", grade="高三", school_year="2026-2027",
                      starts_on=date(2026, 8, 1), ends_on=date(2027, 7, 31))
    db.add(cycle)
    db.flush()
    for cls, student_id, math, english, maximum in (
        (a, ids[0], 90, 80, 100), (a, ids[1], 90, 80, 100), (a, ids[2], 60, None, 100),
        (b, ids[3], 95, 90, 100), (b, ids[4], 90, 80, 100), (c, ids[5], 100, 100, 100),
        (d, ids[6], 100, 90, 150), (a, ids[7], 0, 0, 100),
    ):
        db.add(ClassStudent(class_id=cls.id, student_id=student_id))
        initial_case(db, cls, student_id, ["数学", "英语"], cycle=cycle)
        for subject_name, score, max_score in (("数学", math, maximum), ("英语", english, 100)):
            if score is not None:
                db.add(WeeklyTestScore(class_id=cls.id, student_id=student_id, subject=subject_name,
                                       score=score, max_score=max_score, exam_month="2026-09",
                                       exam_date=date(2026, 10, 1), recorded_by=seed_users["teacher1"]))
    db.commit()
    return classes, ids


def test_totals_ties_missing_zero_and_comparable_scope(client, auth, exam_roster):
    classes, ids = exam_roster
    response = client.get(API + "/overview", headers=auth("admin"))
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["exam_month"] == "2026-09"  # 考试日期跨月不改变月份口径。
    assert data["total"] == 8
    rows = {row["student_id"]: row for row in data["items"]}
    for student_id in ids[:2]:
        row = rows[student_id]
        assert row["total_score"] == 170 and row["total_max_score"] == 200
        assert row["class_rank"] == 1 and row["scope_rank"] == 2
        assert row["scores"]["数学"]["class_rank"] == 1
        assert row["scores"]["数学"]["scope_rank"] == 2
    missing = rows[ids[2]]
    assert missing["total_score"] == 60 and missing["missing_subjects"] == ["英语"]
    assert missing["complete"] is False and missing["class_rank"] is None and missing["scope_rank"] is None
    assert rows[ids[3]]["scope_rank"] == 1
    assert rows[ids[5]]["scope_rank"] == 1  # 高二独立比较。
    assert rows[ids[6]]["scope_rank"] == 1  # 数学满分150独立比较。
    assert rows[ids[7]]["total_score"] == 0 and rows[ids[7]]["class_rank"] == 3
    assert rows[ids[7]]["scores"]["数学"]["class_rank"] == 4
    deyu = client.get(API + "/overview", headers=auth("deyu1")).json()
    assert deyu == data


def test_filters_and_pagination_preserve_ranking_population(client, auth, exam_roster):
    classes, ids = exam_roster
    headers = auth("admin")
    data = client.get(API + "/overview", headers=headers,
                      params=dict(class_id=classes[0].id, limit=1, offset=1)).json()
    assert data["total"] == 4 and len(data["items"]) == 1
    assert data["items"][0]["student_id"] == ids[1]
    assert data["items"][0]["class_rank"] == 1 and data["items"][0]["scope_rank"] == 2
    search = client.get(API + "/overview", headers=headers, params={"keyword": "李四"}).json()
    assert search["total"] == 1 and search["items"][0]["scope_rank"] == 2
    math = client.get(API + "/overview", headers=headers, params={"class_id": classes[0].id, "subject": "数学"}).json()
    assert [row["scores"]["数学"]["score"] for row in math["items"]] == [90, 90, 60, 0]
    empty = client.get(API + "/overview", headers=headers, params={"exam_month": "2026-08"}).json()
    assert empty["total"] == 0 and empty["items"] == []


def test_subject_and_dual_role_scopes_cannot_reveal_other_scores(client, auth, exam_roster):
    classes, ids = exam_roster
    headers = auth("exam_subject")
    data = client.get(API + "/overview", headers=headers).json()
    assert data["total"] == 6 and data["has_total_access"] is False
    for row in data["items"]:
        assert set(row["scores"]) == ({"数学"} if row["class_id"] == classes[0].id else {"英语"})
        assert row["total_score"] is None and row["total_max_score"] is None
        assert row["class_rank"] is None and row["scope_rank"] is None
    for params in ({"class_id": classes[2].id}, {"class_id": classes[0].id, "subject": "英语"}):
        assert client.get(API + "/overview", headers=headers, params=params).json()["items"] == []
    head = client.get(API + "/overview", headers=auth("teacher1")).json()
    assert head["total"] == 4 and all(row["can_view_total"] for row in head["items"])
    dual = client.get(API + "/overview", headers=auth("teacher2")).json()
    for row in dual["items"]:
        if row["class_id"] == classes[0].id:
            assert set(row["scores"]) == {"数学"} and row["total_score"] is None
        else:
            assert row["can_view_total"] is True
    # 旧 teacher 身份在其他班只任教学科时，列表、趋势、汇总同样不能越权。
    for endpoint, params in (("", {"class_id": classes[0].id, "subject": "英语"}),
                             ("/trend", {"student_id": ids[0], "subject": "英语"}),
                             ("/class-summary", {"class_id": classes[0].id, "subject": "英语"})):
        assert client.get(API + endpoint, headers=auth("teacher2"), params=params).json() == []
    assert client.post(API, headers=auth("teacher2"), json=dict(class_id=classes[0].id, student_id=ids[0],
                       subject="英语", exam_date="2026-09-30", score=99, max_score=100)).status_code == 403


@pytest.mark.parametrize("viewer", ["student1", "parent1"])
def test_family_cannot_read_ranking_roster(client, auth, exam_roster, viewer):
    assert client.get(API + "/overview", headers=auth(viewer)).status_code == 403


def test_each_student_uses_own_initial_subjects_not_class_union(client, auth, db, exam_roster, seed_users):
    classes, ids = exam_roster
    case = db.query(StudentCase).filter_by(class_id=classes[0].id, student_id=ids[2]).one()
    db.query(SubjectPlan).filter_by(student_case_id=case.id, subject="英语").delete()
    # 保留误录历史数据，成绩册仅统计本人初评中的数学。
    db.add(WeeklyTestScore(class_id=classes[0].id, student_id=ids[2], subject="英语", score=100,
                           max_score=100, exam_month="2026-09", exam_date=date(2026, 9, 30),
                           recorded_by=seed_users["teacher1"]))
    db.add(SubjectPlan(student_case_id=case.id, subject="德育", teacher_id=classes[0].teacher_id))
    db.commit()
    result = client.get(API + "/overview", headers=auth("admin")).json()
    row = next(row for row in result["items"] if row["student_id"] == ids[2])
    assert row["study_subjects"] == ["数学"] and set(row["scores"]) == {"数学"}
    assert row["total_score"] == 60 and row["total_max_score"] == 100
    assert row["complete"] and row["missing_subjects"] == []
    assert row["class_rank"] == 1 and row["scope_rank"] == 1
    assert "德育" not in result["subjects"]
    assert db.query(WeeklyTestScore).filter_by(student_id=ids[2], subject="英语").one().score == 100
    subject_view = client.get(API + "/overview", headers=auth("exam_subject")).json()
    assert all(set(row["study_subjects"]) <= set(row["scores"]) for row in subject_view["items"])


def test_initial_subject_with_no_class_score_still_controls_total_rank(client, auth, db, exam_roster):
    classes, ids = exam_roster
    case = db.query(StudentCase).filter_by(class_id=classes[0].id, student_id=ids[0]).one()
    db.add(SubjectPlan(student_case_id=case.id, subject="地理", teacher_id=classes[0].teacher_id))
    db.commit()
    result = client.get(API + "/overview", headers=auth("admin")).json()
    row = next(row for row in result["items"] if row["student_id"] == ids[0])
    assert "地理" in result["subjects"] and "地理" in row["study_subjects"]
    assert "地理" not in row["scores"] and row["total_score"] == 170
    assert row["complete"] is False and row["class_rank"] is None


def test_missing_initial_never_falls_back_to_class_or_archived_case(client, auth, db, exam_roster):
    classes, ids = exam_roster
    case = db.query(StudentCase).filter_by(class_id=classes[0].id, student_id=ids[0]).one()
    db.query(SubjectPlan).filter_by(student_case_id=case.id).delete()
    initial_case(db, classes[0], ids[0], ["数学", "英语"], status="archived")
    initial_case(db, classes[1], ids[0], ["数学", "英语"])
    db.commit()
    result = client.get(API + "/overview", headers=auth("admin")).json()
    row = next(row for row in result["items"] if row["student_id"] == ids[0])
    assert row["study_subjects"] == [] and row["scores"] == {}
    assert row["total_score"] is None and row["class_rank"] is None and row["complete"] is False
    subject_view = client.get(API + "/overview", headers=auth("exam_subject")).json()
    assert ids[0] not in {row["student_id"] for row in subject_view["items"]}


def test_aggregate_does_not_truncate_at_score_list_limit(client, auth, db, seed_users):
    cls = Class(name="大班", grade="高三", teacher_id=seed_users["teacher1"])
    db.add(cls)
    db.flush()
    cycle = CaseCycle(name="大班初评", grade="高三", school_year="2026-2027",
                      starts_on=date(2026, 8, 1), ends_on=date(2027, 7, 31))
    db.add(cycle)
    db.flush()
    password = hash_password("test123456")
    for i in range(205):
        student = User(username=f"many_{i}", name=f"学生{i}", role="student", password_hash=password)
        db.add(student)
        db.flush()
        initial_case(db, cls, student.id, ["数学"], cycle=cycle)
        db.add(WeeklyTestScore(class_id=cls.id, student_id=student.id, subject="数学", score=i,
                               max_score=300, exam_month="2026-09", exam_date=date(2026, 9, 30),
                               recorded_by=seed_users["teacher1"]))
    db.commit()
    response = client.get(API + "/overview", headers=auth("teacher1"), params={"limit": 5, "offset": 200})
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["total"] == 205 and len(data["items"]) == 5
    assert [row["class_rank"] for row in data["items"]] == [201, 202, 203, 204, 205]


def test_export_overview_workbook_uses_same_scope_and_no_page_limit(client, auth, exam_roster):
    from io import BytesIO

    from openpyxl import load_workbook

    response = client.get(API + "/overview/export", headers=auth("admin"))
    assert response.status_code == 200, response.text
    assert "spreadsheetml" in response.headers["content-type"]
    assert "filename*=UTF-8''" in response.headers["content-disposition"]
    ws = load_workbook(BytesIO(response.content)).active
    headers = [cell.value for cell in ws[2]]
    assert headers[:3] == ["学生姓名", "班级", "年级"]
    assert {"总分", "总分班内排名", "数学得分", "数学班内排名", "英语范围排名"} <= set(headers)
    rows = list(ws.iter_rows(min_row=3, values_only=True))
    assert len(rows) == 8
    math_index = headers.index("数学得分")
    rank_index = headers.index("数学班内排名")
    assert sorted(row[math_index] for row in rows) == [0, 60, 90, 90, 90, 95, 100, 100]
    assert any(row[rank_index] == 1 for row in rows)

    subject_ws = load_workbook(BytesIO(client.get(API + "/overview/export", headers=auth("exam_subject")).content)).active
    subject_headers = [cell.value for cell in subject_ws[2]]
    assert "总分" not in subject_headers and "数学得分" in subject_headers and "英语得分" in subject_headers

    assert client.get(API + "/overview/export", headers=auth("student1")).status_code == 403
    assert client.get(API + "/overview/export", headers=auth("admin"),
                      params={"exam_month": "2026-08"}).status_code == 404
